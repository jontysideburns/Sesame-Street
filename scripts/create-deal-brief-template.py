"""Generate a deal brief template — one deal at a time, enough for full topsheet synthesis."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
thin = Side(style="thin", color="CCCCCC")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
hdr_fill = PatternFill("solid", fgColor="D6E4F0")
hdr_font = Font(bold=True, size=10, name="Arial")
note_font = Font(italic=True, size=9, color="666666", name="Arial")
blue = Font(size=10, color="0000FF", name="Arial")
body = Font(size=10, name="Arial")
section_fill = PatternFill("solid", fgColor="1F6FA5")
section_font = Font(bold=True, size=10, color="FFFFFF", name="Arial")
req_font = Font(bold=True, size=10, color="CC0000", name="Arial")

ws = wb.active
ws.title = "Deal Brief"
ws.column_dimensions["A"].width = 35
ws.column_dimensions["B"].width = 55
ws.column_dimensions["C"].width = 45

# Title
ws.merge_cells("A1:C1")
ws["A1"] = "DEAL BRIEF — Fill in column B. Claude generates the full topsheet + management case forecast from this."
ws["A1"].font = Font(bold=True, size=11, name="Arial")
ws["A1"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[1].height = 35

row = 2

def add_section(title):
    global row
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    cell = ws.cell(row=row, column=1, value=title)
    cell.font = section_font
    cell.fill = section_fill
    for c in range(1, 4):
        ws.cell(row=row, column=c).border = border

def add_field(label, guidance, required=False):
    global row
    row += 1
    ws.cell(row=row, column=1, value=label).font = req_font if required else body
    ws.cell(row=row, column=2).font = blue
    ws.cell(row=row, column=3, value=guidance).font = note_font
    for c in range(1, 4):
        ws.cell(row=row, column=c).border = border
        ws.cell(row=row, column=c).alignment = Alignment(vertical="top", wrap_text=True)

# Headers
row += 1
for i, h in enumerate(["Field", "Your Input", "Guidance"], 1):
    ws.cell(row=row, column=i, value=h)
for c in range(1, 4):
    ws.cell(row=row, column=c).font = hdr_font
    ws.cell(row=row, column=c).fill = hdr_fill
    ws.cell(row=row, column=c).border = border

# ── Section 1: Identity ──
add_section("1. DEAL IDENTITY")
add_field("Deal Name *", "Full name as it should appear in the system", True)
add_field("Borrower Legal Name *", "SPV or borrower entity name", True)
add_field("Sector *", "Wind Farm, Data Center, Port, Airport, Toll Road, Social Infra, Real Estate, Clean Tech Hub", True)
add_field("Sub-Sector", "e.g. Offshore Wind, Hyperscale Colocation, Container Terminal")
add_field("Country *", "ISO 2-letter code (GB, US, DE, etc.)", True)
add_field("Region *", "EMEA, North America, APAC", True)
add_field("Currency *", "GBP, USD, EUR", True)
add_field("Sponsor / Equity Owner", "Who owns the equity?")
add_field("Phase *", "construction, ramp_up, operational, refinancing", True)

# ── Section 2: Structure ──
add_section("2. CAPITAL STRUCTURE & TERMS")
add_field("Total Facility Size (M) *", "Total committed debt across all tranches, in millions", True)
add_field("Our Exposure (M) *", "Our share / holding, in millions", True)
add_field("Security Ranking *", "Senior Secured, Senior Unsecured, Mezzanine, Holdco, Majority Holdco, Minority Holdco", True)
add_field("Number of Tranches", "e.g. 1 (single note), 2 (senior + mezz), 3 (senior + capex + SHL)")
add_field("Instrument Type *", "loan, note, bond, frn, il_bond", True)
add_field("Margin (bps) *", "e.g. 200, 350", True)
add_field("Base Rate", "SONIA, SOFR, GILT, Euribor, or 'fixed'")
add_field("Fixed or Floating *", "fixed or floating", True)
add_field("Repayment Type", "sculpted, amortising, bullet")
add_field("Origination Date", "YYYY-MM-DD (when was this deal done?)")
add_field("Maturity Date *", "YYYY-MM-DD", True)
add_field("Concession / Contract Expiry", "If applicable (YYYY-MM-DD)")
add_field("Governing Law", "English, New York, etc.")

# ── Section 3: Revenue Model ──
add_section("3. REVENUE MODEL (2-3 sentences)")
add_field("Revenue Description *", "What drives revenue? Contracted or merchant? Who is the offtaker? What is the contract?", True)
add_field("Contracted Revenue %", "0-100 (what % is under contract?)")
add_field("Key Revenue Contract", "e.g. CfD with LCCC, PPA with Shell, Availability Payment from NHS")
add_field("Contract Expiry", "When does the main revenue contract end?")
add_field("Revenue Risk Code", "e.g. P1-V1-D2 (I can help assign this)")

# ── Section 4: Cost Structure ──
add_section("4. COST STRUCTURE (2-3 sentences)")
add_field("Cost Description *", "What are the big costs? O&M contract? Operator? Index-linked?", True)
add_field("EBITDA Margin %", "Approximate (e.g. 85% for wind, 50% for data centre)")
add_field("Key Cost Contracts", "e.g. O&M with Siemens Gamesa, FM with Bouygues")

# ── Section 5: Current Financials ──
add_section("5. CURRENT FINANCIALS (latest reporting period)")
add_field("Latest Period", "e.g. H2 2025, Q2 2026")
add_field("Revenue (M, annual) *", "In deal currency, millions", True)
add_field("EBITDA (M, annual) *", "In deal currency, millions", True)
add_field("Current DSCR *", "e.g. 1.35x", True)
add_field("Net Debt (M)", "Outstanding debt minus cash")
add_field("Cash Balance (M)", "")
add_field("Capex (M, annual)", "Maintenance capex")

# ── Section 6: Ratings ──
add_section("6. RATINGS")
add_field("Moodys Rating", "e.g. Baa2, Ba1, or n/a")
add_field("S&P Rating", "e.g. BBB, BB+, or n/a")
add_field("Fitch Rating", "e.g. BBB, or n/a")
add_field("Internal Credit Score *", "Your internal view (e.g. BBB, BB+)", True)

# ── Section 7: Key KPIs ──
add_section("7. KEY SECTOR KPIs (3-5 metrics)")
add_field("KPI 1 (name + current value)", "e.g. Capacity Factor: 42%, Occupancy: 85%, AADT: 35,000")
add_field("KPI 2 (name + current value)", "")
add_field("KPI 3 (name + current value)", "")
add_field("KPI 4 (name + current value)", "")
add_field("KPI 5 (name + current value)", "")

# ── Section 8: Performance & Direction ──
add_section("8. PERFORMANCE & DIRECTION OF TRAVEL")
add_field("Performance Narrative *", "2-3 sentences: How is the deal performing? Any concerns?", True)
add_field("Direction of Travel *", "improving, stable, gradually deteriorating, rapidly worsening", True)
add_field("Current Grade *", "1-Outperforming, 2-In Line, 3-Underperforming, 4-Stressed", True)
add_field("On Watchlist?", "Y or N")

# ── Section 9: Stress Scenario ──
add_section("9. STRESS SCENARIO (for IC memo KPI targets)")
add_field("Main Downside Risk *", "1-2 sentences: What could go wrong? What does the stress case look like?", True)
add_field("Stress DSCR", "What would DSCR fall to under stress? e.g. 1.05x")

# ── Section 10: Anything Else ──
add_section("10. ANYTHING ELSE")
add_field("Additional Notes", "Anything else relevant: upcoming refinancing, regulatory review, key person risk, etc.")
add_field("Key Counterparties", "List 2-3 important counterparties if known (offtaker, operator, insurer)")
add_field("Reserve Accounts", "Any specifics: DSRA sizing, LC-backed, partially funded?")

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "deal-brief-template.xlsx")
wb.save(out)
print(f"Saved to {out}")
