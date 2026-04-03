"""Generate a lightweight deal onboarding template for quickly seeding test deals."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
thin = Side(style="thin", color="CCCCCC")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
hdr_fill = PatternFill("solid", fgColor="D6E4F0")
hdr_font = Font(bold=True, size=10, name="Arial")
note_font = Font(italic=True, size=9, color="666666", name="Arial")
blue = Font(size=10, color="0000FF", name="Arial")
body = Font(size=10, name="Arial")
req = Font(bold=True, size=10, color="CC0000", name="Arial")
section_fill = PatternFill("solid", fgColor="1F6FA5")
section_font = Font(bold=True, size=10, color="FFFFFF", name="Arial")

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="top")

def section_row(ws, row, cols, text):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = section_font
    cell.fill = section_fill
    cell.alignment = Alignment(vertical="center")
    for c in range(1, cols + 1):
        ws.cell(row=row, column=c).border = border

ws = wb.active
ws.title = "Deals"

# Instructions
ws.merge_cells("A1:F1")
ws["A1"] = "QUICK DEAL ONBOARDING - Fill in one row per deal. Provide essentials + directional guidance. Claude fills in the rest."
ws["A1"].font = Font(bold=True, size=11, name="Arial")
ws["A1"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[1].height = 35

ws.merge_cells("A2:F2")
ws["A2"] = 'For "Performance" and "Direction": use plain English (e.g. "strong performer", "under pressure from rising costs", "construction 6 months behind"). Claude will generate realistic financials, covenants, counterparties, and KPIs.'
ws["A2"].font = note_font
ws["A2"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 30

# Headers
row = 3
headers = [
    # Core identity
    ("Deal Name *", 22, "e.g. Thames Water WBS"),
    ("Sector *", 18, "Wind Farm, Data Center, Port, Airport, Toll Road, Social Infra, Real Estate, Clean Tech Hub"),
    ("Sub-Sector", 18, "e.g. Offshore Wind, Hyperscale Colocation"),
    ("Country *", 8, "ISO 2-letter"),
    ("Region *", 12, "EMEA, North America, APAC"),
    ("Currency *", 8, "GBP, USD, EUR"),
    # Structure
    ("Borrower Name *", 22, "Legal entity name"),
    ("Deal Type *", 16, "Project Finance, Acquisition Finance"),
    ("Phase *", 14, "construction, ramp_up, operational, refinancing"),
    ("Security Ranking *", 18, "Senior Secured, Mezzanine, Holdco, etc."),
    # Economics
    ("Total Facility (M) *", 14, "In deal currency, millions"),
    ("Our Exposure (M) *", 14, "Our share, millions"),
    ("Instrument Type *", 14, "loan, note, bond, frn"),
    ("Margin (bps) *", 10, "e.g. 200"),
    ("Base Rate", 10, "SONIA, SOFR, GILT, fixed"),
    ("Maturity Date *", 12, "YYYY-MM-DD"),
    ("WAL (years)", 10, "e.g. 7.5"),
    # Ratings
    ("Moodys", 8, "e.g. Baa2"),
    ("S&P", 8, "e.g. BBB"),
    ("Fitch", 8, "e.g. BBB"),
    ("Internal Score *", 8, "If no external"),
    # Revenue
    ("Contracted Rev %", 10, "0-100"),
    ("Revenue Risk Code", 12, "e.g. P1-V1-D2"),
    # Performance guidance
    ("Current Revenue (M)", 14, "Annual, millions"),
    ("Current EBITDA (M)", 14, "Annual, millions"),
    ("Current DSCR", 10, "e.g. 1.35"),
    ("Performance *", 25, "Plain English: strong, adequate, under pressure, stressed"),
    ("Direction *", 25, "Plain English: improving, stable, gradually deteriorating, rapidly worsening"),
    ("Key Issues", 35, "Any specific concerns: lease-up slow, cost overruns, counterparty risk, etc."),
    # Watchlist
    ("Watchlist", 8, "Y or N"),
    ("Grade", 14, "1-Outperforming, 2-In Line, 3-Underperforming, 4-Stressed"),
]

for i, (label, width, note) in enumerate(headers, 1):
    ws.cell(row=row, column=i, value=label)
    ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = width
style_header(ws, row, len(headers))

# Notes row
for i, (_, _, note) in enumerate(headers, 1):
    cell = ws.cell(row=row + 1, column=i, value=note)
    cell.font = note_font
    cell.border = border
    cell.alignment = Alignment(wrap_text=True, vertical="top")

# Example row
example = [
    "North Sea OWF", "Wind Farm", "Offshore Wind", "GB", "EMEA", "GBP",
    "North Sea OWF Ltd", "Project Finance", "operational", "Senior Secured",
    145, 100, "note", 200, "GILT", "2044-12-31", 13,
    "", "", "", "Baa2",
    100, "P1-V1-D2",
    20, 17, 1.23, "Adequate performer, CfD-backed", "Stable, slight MRA shortfall", "MRA partially funded",
    "N", "2 - In Line",
]
for i, v in enumerate(example, 1):
    cell = ws.cell(row=row + 2, column=i, value=v)
    cell.font = blue
    cell.border = border

# 15 empty rows for new deals
for r in range(row + 3, row + 18):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = border
        cell.font = blue

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "quick-deal-onboarding.xlsx")
wb.save(out)
print(f"Saved to {out}")
