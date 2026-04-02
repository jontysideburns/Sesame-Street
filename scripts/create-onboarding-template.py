"""Generate the deal onboarding Excel template."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
thin = Side(style="thin", color="CCCCCC")
border = Border(top=thin, bottom=thin, left=thin, right=thin)
hdr_fill = PatternFill("solid", fgColor="D6E4F0")
hdr_font = Font(bold=True, size=10, name="Arial")
note_font = Font(italic=True, size=9, color="666666", name="Arial")
req_font = Font(bold=True, size=10, color="CC0000", name="Arial")
body_font = Font(size=10, name="Arial")
blue_font = Font(size=10, color="0000FF", name="Arial")


def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.border = border
        cell.alignment = Alignment(wrap_text=True, vertical="top")


def style_rows(ws, start, end, cols, input_col=2):
    for r in range(start, end + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.font = blue_font if c == input_col else body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def note_row(ws, row, notes):
    for i, n in enumerate(notes, 1):
        cell = ws.cell(row=row, column=i, value=n)
        cell.font = note_font
        cell.border = border


# ── Sheet 1: Deal Identity ──
ws = wb.active
ws.title = "Deal Identity"
ws.merge_cells("A1:C1")
ws["A1"] = "DEAL ONBOARDING TEMPLATE - Fill in column B. Required fields marked with *."
ws["A1"].font = Font(bold=True, size=11, name="Arial")
ws["A1"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[1].height = 30

for i, h in enumerate(["Field", "Value", "Notes / Guidance"], 1):
    ws.cell(row=2, column=i, value=h)
style_header(ws, 2, 3)

fields = [
    ("Investment Name *", "", "Full deal name"),
    ("Borrower Legal Name *", "", "Registered company name"),
    ("Borrower Trading Name", "", "If different from legal name"),
    ("Borrower LEI", "", "20-character Legal Entity Identifier"),
    ("Borrower Jurisdiction *", "", "ISO 2-letter (e.g. US, GB, NL)"),
    ("Borrower Registered Address", "", "Full registered office address"),
    ("Sector *", "", "Data Center, Wind Farm, Port, Airport, Toll Road, Social Infrastructure, Real Estate, Clean Tech Hub"),
    ("Sub-Sector", "", "e.g. Hyperscale Colocation, Offshore Wind"),
    ("Deal Type *", "", "Project Finance, Acquisition Finance, Development Finance"),
    ("Phase *", "", "construction, ramp_up, operational, refinancing"),
    ("Sponsor Name", "", "Lead sponsor / equity investor"),
    ("Sponsor Fund", "", "Fund vehicle name"),
    ("Region *", "", "North America, EMEA, APAC, LATAM, Middle East Africa"),
    ("Country *", "", "ISO 2-letter (e.g. US, GB, DE, SG)"),
    ("Currency *", "", "ISO 3-letter (e.g. USD, GBP, EUR)"),
    ("Origination Date *", "", "YYYY-MM-DD"),
    ("Commitment Date", "", "YYYY-MM-DD"),
    ("First Drawdown Date", "", "YYYY-MM-DD"),
    ("COD Date", "", "Commercial Operations Date"),
    ("Maturity Date *", "", "Final debt maturity"),
    ("Weighted Average Life", "", "Years (e.g. 7.5)"),
    ("Concession Expiry Date", "", "For concession-based deals"),
    ("Fiscal Year End Month *", "", "1-12 (e.g. 12 for December)"),
    ("Governing Law", "", "e.g. English, New York"),
    ("Security Ranking *", "", "Senior Secured, Senior Unsecured, Second Lien, Mezzanine, Subordinated, Holdco, Majority Holdco, Minority Holdco"),
    ("Revenue Risk Code", "", "e.g. P2-V4-D2"),
    ("Contracted Revenue %", "", "0-100"),
    ("Merchant Revenue %", "", "0-100"),
    ("Duration Coverage %", "", "Contract life / debt term x 100"),
    ("Moodys Rating", "", "e.g. Baa2, Ba1"),
    ("S&P Rating", "", "e.g. BBB, BB+"),
    ("Fitch Rating", "", "e.g. BBB, BB+"),
    ("Internal Credit Score *", "", "Required if no external rating"),
]
for i, (field, val, note) in enumerate(fields, 3):
    ws.cell(row=i, column=1, value=field)
    ws.cell(row=i, column=2, value=val)
    ws.cell(row=i, column=3, value=note)
style_rows(ws, 3, 3 + len(fields) - 1, 3)
for i, (field, _, _) in enumerate(fields, 3):
    if "*" in field:
        ws.cell(row=i, column=1).font = req_font

ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 35
ws.column_dimensions["C"].width = 55

# ── Sheet 2: Capital Structure ──
ws2 = wb.create_sheet("Capital Structure")
caps = ["Instrument Name", "Type", "Format", "Security Ranking", "Pari-Passu Group",
        "Committed Amount", "Drawn Amount", "Currency", "Margin (bps)", "Base Rate",
        "Interest Type", "Maturity Date", "Repayment Type", "Our Holding Amount",
        "Our Holding %", "DSRA Months", "Status"]
for i, h in enumerate(caps, 1):
    ws2.cell(row=1, column=i, value=h)
style_header(ws2, 1, len(caps))
note_row(ws2, 2, [
    "", "senior_term, senior_rcf, capex_facility, mezzanine, bond, note, frn, shl",
    "loan, bond, note, frn, il_bond", "Senior, Junior, Holdco", "A, B (same = pari-passu)",
    "", "", "USD/GBP/EUR", "", "SONIA, SOFR, Euribor",
    "fixed, floating", "YYYY-MM-DD", "bullet, amortising, sculpted",
    "", "%", "Integer", "active"
])
for r in range(3, 8):
    for c in range(1, len(caps) + 1):
        ws2.cell(row=r, column=c).border = border
        ws2.cell(row=r, column=c).font = blue_font
for c in range(1, len(caps) + 1):
    ws2.column_dimensions[ws2.cell(row=1, column=c).column_letter].width = max(14, len(caps[c - 1]) + 2)

# ── Sheet 3: Reserve Accounts ──
ws3 = wb.create_sheet("Reserve Accounts")
res = ["Account Name", "Type", "Sizing Basis", "Required Balance", "Current Balance",
       "Cash Amount", "LC Amount", "LC Provider", "PCG Amount", "PCG Provider",
       "Funded Status", "Currency"]
for i, h in enumerate(res, 1):
    ws3.cell(row=1, column=i, value=h)
style_header(ws3, 1, len(res))
note_row(ws3, 2, [
    "", "dsra, mra, capex_reserve, lifecycle_reserve, liquidity_facility",
    "e.g. 6 months senior DS", "", "", "", "", "", "", "",
    "fully_funded, partially_funded, unfunded", "USD/GBP/EUR"
])
for r in range(3, 7):
    for c in range(1, len(res) + 1):
        ws3.cell(row=r, column=c).border = border
        ws3.cell(row=r, column=c).font = blue_font
for c in range(1, len(res) + 1):
    ws3.column_dimensions[ws3.cell(row=1, column=c).column_letter].width = max(14, len(res[c - 1]) + 2)

# ── Sheet 4: Counterparties ──
ws4 = wb.create_sheet("Counterparties")
cps = ["Name", "Type", "Credit Rating", "Contract Value", "Contract Expiry",
       "Replacement Risk", "Notes"]
for i, h in enumerate(cps, 1):
    ws4.cell(row=1, column=i, value=h)
style_header(ws4, 1, len(cps))
note_row(ws4, 2, [
    "", "offtaker, contractor, operator, guarantor, insurer, auditor, facility_agent",
    "e.g. A+/Stable", "", "YYYY-MM-DD", "low, medium, high, critical", ""
])
for r in range(3, 8):
    for c in range(1, len(cps) + 1):
        ws4.cell(row=r, column=c).border = border
        ws4.cell(row=r, column=c).font = blue_font
ws4.column_dimensions["A"].width = 30
ws4.column_dimensions["B"].width = 22
ws4.column_dimensions["G"].width = 35

# ── Sheet 5: Hedging ──
ws5 = wb.create_sheet("Hedging")
hdg = ["Hedge Type", "Notional", "% of Debt", "Fixed Rate", "Strike",
       "Counterparty", "Counterparty Rating", "Mark to Market", "Maturity Date"]
for i, h in enumerate(hdg, 1):
    ws5.cell(row=1, column=i, value=h)
style_header(ws5, 1, len(hdg))
note_row(ws5, 2, [
    "interest_rate_swap, cap, floor, fx_forward, inflation_swap",
    "", "%", "%", "%", "", "e.g. A+/Stable", "", "YYYY-MM-DD"
])
for r in range(3, 6):
    for c in range(1, len(hdg) + 1):
        ws5.cell(row=r, column=c).border = border
        ws5.cell(row=r, column=c).font = blue_font
for c in range(1, len(hdg) + 1):
    ws5.column_dimensions[ws5.cell(row=1, column=c).column_letter].width = max(14, len(hdg[c - 1]) + 2)

# ── Sheet 6: Jurisdiction Splits ──
ws6 = wb.create_sheet("Jurisdiction Splits")
jur = ["Country Code", "Country Name", "Activity %", "Activity Type", "Is Primary"]
for i, h in enumerate(jur, 1):
    ws6.cell(row=1, column=i, value=h)
style_header(ws6, 1, len(jur))
note_row(ws6, 2, ["ISO 2-letter", "", "Must sum to 100%", "revenue", "TRUE or FALSE"])
for r in range(3, 7):
    for c in range(1, len(jur) + 1):
        ws6.cell(row=r, column=c).border = border
        ws6.cell(row=r, column=c).font = blue_font
for c in range(1, len(jur) + 1):
    ws6.column_dimensions[ws6.cell(row=1, column=c).column_letter].width = 18

# ── Sheet 7: Key Metrics ──
ws7 = wb.create_sheet("Key Metrics")
ws7.merge_cells("A1:C1")
ws7["A1"] = "Current period financial snapshot"
ws7["A1"].font = Font(bold=True, size=11, name="Arial")
for i, h in enumerate(["Metric", "Value", "Notes"], 1):
    ws7.cell(row=2, column=i, value=h)
style_header(ws7, 2, 3)
metrics = [
    ("Revenue", ""), ("EBITDA", ""), ("CFADS", ""), ("Debt Service", ""),
    ("Net Debt", ""), ("Cash", ""), ("Senior DSCR", "e.g. 1.35"),
    ("Net Debt / EBITDA", "e.g. 5.2"), ("Latest Period Label", "e.g. Q2 2026"),
    ("Latest Period End Date", "YYYY-MM-DD"),
]
for i, (m, n) in enumerate(metrics, 3):
    ws7.cell(row=i, column=1, value=m)
    ws7.cell(row=i, column=3, value=n)
style_rows(ws7, 3, 12, 3)
ws7.column_dimensions["A"].width = 25
ws7.column_dimensions["B"].width = 20
ws7.column_dimensions["C"].width = 25

# ── Sheet 8: KPI Targets ──
ws8 = wb.create_sheet("KPI Targets")
ws8.merge_cells("A1:F1")
ws8["A1"] = "IC Memo KPI targets - frozen at ingestion. Base case and stress case expectations."
ws8["A1"].font = Font(bold=True, size=11, name="Arial")
ws8["A1"].alignment = Alignment(wrap_text=True)
kpi = ["KPI Name", "Base Case Value", "Stress Case Value", "Direction", "Unit", "Source"]
for i, h in enumerate(kpi, 1):
    ws8.cell(row=2, column=i, value=h)
style_header(ws8, 2, len(kpi))
note_row(ws8, 3, [
    "e.g. Leased Capacity (%)", "", "",
    "higher_is_better or lower_is_better",
    "percentage, currency, count, ratio, years",
    "ic_memo, business_plan"
])
for r in range(4, 14):
    for c in range(1, len(kpi) + 1):
        ws8.cell(row=r, column=c).border = border
        ws8.cell(row=r, column=c).font = blue_font
for c in range(1, len(kpi) + 1):
    ws8.column_dimensions[ws8.cell(row=2, column=c).column_letter].width = max(18, len(kpi[c - 1]) + 4)

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs", "deal-onboarding-template.xlsx")
wb.save(out)
print(f"Saved to {out}")
